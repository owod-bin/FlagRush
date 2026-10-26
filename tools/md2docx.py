# -*- coding: utf-8 -*-
"""把 使用说明.md 转成 Word（.docx）

为什么自己写：项目里没有 pandoc，而 python-docx 直接贴 markdown 会把
`反引号`、**星号**、表格竖线原样留在文档里（旧版 docx 就是这样）。
这里做一个够用的子集转换：

  # / ## / ###       → Heading 1/2/3
  ```代码块```        → Consolas + 浅灰底
  | 表格 |            → Table Grid，首行加粗
  - / * 列表           → List Bullet（按缩进分级）
  1. 列表              → List Number
  > 引用               → 斜体 + 缩进
  **粗体** `代码` *斜体* → 行内 run

用法：
    python tools\\md2docx.py 使用说明.md 使用说明.docx
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

CN_FONT = "微软雅黑"
CODE_FONT = "Consolas"


def _set_cn(run, font=CN_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), font)


def _shade(paragraph, color="F2F2F2"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    ppr.append(shd)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_inline(p, text, base_size=10.5, base_bold=False, base_color=None):
    """把一行里的 **粗体** / `代码` / *斜体* 拆成 run（支持 **`代码`** 嵌套）"""
    for seg in INLINE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            # 递归处理 **`code`** 这种「粗体里套代码」
            add_inline(p, seg[2:-2], base_size, True, base_color)
            continue
        if seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            r = p.add_run(seg[1:-1])
            r.font.name = CODE_FONT
            r.font.size = Pt(base_size - 1)
            r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            r.bold = base_bold
            _set_cn(r, CODE_FONT)
            continue
        elif seg.startswith("*") and seg.endswith("*") and len(seg) > 2:
            r = p.add_run(seg[1:-1])
            r.italic = True
        else:
            r = p.add_run(seg)
        r.font.size = Pt(base_size)
        r.bold = r.bold or base_bold
        if base_color:
            r.font.color.rgb = base_color
        _set_cn(r)


def _merge_bold_lines(lines):
    """把跨行的 **粗体** 合成一行。

    markdown 里长句常写成：
        **这意味着：把整个文件夹拷到任意电脑上，
        不用装 Python、不用联网。**
    不合并的话，两个 ** 各留一半在文档里。
    """
    out, buf, in_code = [], None, False
    for ln in lines:
        if ln.strip().startswith("```"):
            in_code = not in_code
            if buf is not None:
                out.append(buf)
                buf = None
            out.append(ln)
            continue
        if in_code:
            out.append(ln)
            continue
        if buf is not None:
            buf = buf.rstrip() + ln.lstrip()
            if buf.count("**") % 2 == 0:
                out.append(buf)
                buf = None
            continue
        if ln.count("**") % 2 == 1:
            buf = ln
            continue
        out.append(ln)
    if buf is not None:
        out.append(buf)
    return out


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(cols):
            txt = row[j] if j < len(row) else ""
            cell = cells[j]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, txt, base_size=9.5, base_bold=(i == 0))
    doc.add_paragraph()


def convert(md_path, docx_path):
    lines = _merge_bold_lines(open(md_path, encoding="utf-8").read().splitlines())
    doc = Document()

    # 正文默认字体
    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    for lvl, size in ((1, 18), (2, 15), (3, 12.5), (4, 11.5)):
        st = doc.styles["Heading %d" % lvl]
        st.font.name = CN_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.rstrip()

        # 代码块
        if s.strip().startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            for code in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.3)
                r = p.add_run(code if code else " ")
                r.font.name = CODE_FONT
                r.font.size = Pt(9)
                _set_cn(r, CODE_FONT)
                _shade(p)
            doc.add_paragraph()
            continue

        # 表格
        if s.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                raw = lines[i].strip().strip("|")
                rows.append([c.strip() for c in raw.split("|")])
                i += 1
            if rows and all(re.match(r"^[\s:-]+$", c or "-") for c in rows[1]):
                rows.pop(1)
            add_table(doc, rows)
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            p = doc.add_heading(level=lvl)
            add_inline(p, m.group(2), base_size={1: 18, 2: 15, 3: 12.5, 4: 11.5}[lvl])
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            i += 1
            continue

        # 分隔线
        if re.match(r"^\s*---+\s*$", s):
            i += 1
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, s.lstrip("> ").strip(), base_color=RGBColor(0x55, 0x55, 0x55))
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)[-*]\s+(.*)$", s)
        if m:
            indent = len(m.group(1)) // 2
            p = doc.add_paragraph(style="List Bullet" if indent == 0 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\s*)\d+[.)]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 空行 / 正文
        if not s.strip():
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, s)
        i += 1

    doc.save(docx_path)
    return docx_path


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    out = convert(sys.argv[1], sys.argv[2])
    print("已生成：%s（%.0f KB）" % (out, os.path.getsize(out) / 1024.0))
    return 0


if __name__ == "__main__":
    sys.exit(main())
